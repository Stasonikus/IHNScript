import os
import sys
import json
from docx import Document
from docx.shared import Pt
from datetime import datetime, timedelta

# ==========================================
# 📁 Project paths (works for .py and .exe)
# ==========================================

def get_base_dir():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    else:
        return os.path.dirname(os.path.abspath(__file__))

BASE_DIR = get_base_dir()

TEMPLATE_DIR = os.path.join(BASE_DIR, "templates")
DATA_DIR = os.path.join(BASE_DIR, "data")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")

CARS_FILE = os.path.join(DATA_DIR, "cars.json")
COUNTER_FILE = os.path.join(DATA_DIR, "contract_counter.json")

os.makedirs(OUTPUT_DIR, exist_ok=True)

# ==========================================
# 🔧 Input validation
# ==========================================
def input_date(prompt):
    while True:
        value = input(f"{prompt} (DD.MM.YYYY, example 05.01.2026): ").strip()
        try:
            datetime.strptime(value, "%d.%m.%Y")
            return value
        except ValueError:
            print("❌ Wrong date format. Use DD.MM.YYYY (example: 05.01.2026)")

def input_float(prompt):
    while True:
        value = input(prompt).strip()
        try:
            number = float(value)
            if number < 0:
                raise ValueError
            return number
        except ValueError:
            print("❌ Enter a valid positive number")

# ==========================================
# 🔧 Replace text in DOCX (keep styles)
# ==========================================
def replace_text_in_doc(doc, replacements):

    DRIVER_BLOCK_KEYS = [
        "{{DRIVER1_BLOCK}}",
        "{{DRIVER2_BLOCK}}",
        "{{DRIVER3_BLOCK}}"
    ]

    def process_paragraph(paragraph):
        if not paragraph.runs:
            return

        full_text = "".join(run.text for run in paragraph.runs)
        new_text = full_text

        for key, value in replacements.items():
            new_text = new_text.replace(key, str(value))

        if new_text != full_text:
            is_driver_block = any(k in full_text for k in DRIVER_BLOCK_KEYS)

            ref_run = paragraph.runs[0]
            font_name = ref_run.font.name
            font_size = ref_run.font.size
            font_bold = ref_run.font.bold
            font_italic = ref_run.font.italic

            if is_driver_block:
                font_name = "Times New Roman"
                font_size = Pt(14)
                font_bold = False
                font_italic = False

            for r in paragraph.runs:
                r._element.getparent().remove(r._element)

            new_run = paragraph.add_run(new_text)
            new_run.font.name = font_name
            new_run.font.size = font_size
            new_run.font.bold = font_bold
            new_run.font.italic = font_italic

    for p in doc.paragraphs:
        process_paragraph(p)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    process_paragraph(p)

# ==========================================
# 🔢 Contract number
# ==========================================
def load_contract_number():
    if not os.path.exists(COUNTER_FILE):
        return 1
    with open(COUNTER_FILE, "r", encoding="utf-8") as f:
        return json.load(f).get("last_number", 0) + 1

def save_contract_number(number):
    with open(COUNTER_FILE, "w", encoding="utf-8") as f:
        json.dump({"last_number": number}, f, ensure_ascii=False, indent=4)

# ==========================================
# 📝 Document generation
# ==========================================
def generate_docs(base_data, client_name, end_date):

    for filename in os.listdir(TEMPLATE_DIR):
        if not filename.endswith(".docx") or filename.startswith("~$"):
            continue

        template_path = os.path.join(TEMPLATE_DIR, filename)
        data = dict(base_data)

        if "contract" in filename.lower():
            out_name = f"contract_{client_name}.docx"
            data["{{RENTAL_END}}"] = end_date

        elif "poa" in filename.lower():
            out_name = f"poa_{client_name}.docx"
            data["{{RENTAL_END}}"] = extend_date(end_date)

        elif "waybill" in filename.lower():
            out_name = f"waybill_{client_name}.docx"
            data["{{RENTAL_END}}"] = extend_date(end_date)

        else:
            out_name = f"{client_name}_{filename}"

        doc = Document(template_path)
        replace_text_in_doc(doc, data)

        out_path = os.path.join(OUTPUT_DIR, out_name)
        doc.save(out_path)
        print(f"✅ File created: {out_path}")

# ==========================================
# 🚗 Load cars
# ==========================================
def load_cars():
    with open(CARS_FILE, "r", encoding="utf-8") as f:
        return json.load(f)

# ==========================================
# 🚗 Choose multiple cars
# ==========================================
def choose_cars(cars):
    print("\nAvailable cars:")
    for i, c in enumerate(cars, 1):
        print(f"{i}. {c['make']} {c['model']} ({c['plate']})")

    while True:
        choice = input("Select car numbers (example: 1 or 1,3): ").strip()
        selected = []

        for c in choice.split(","):
            if c.strip().isdigit():
                idx = int(c.strip())
                if 1 <= idx <= len(cars):
                    selected.append(cars[idx - 1])

        if selected:
            return selected
        else:
            print("❌ Invalid selection. Try again.")

# ==========================================
# 📆 Dates
# ==========================================
def calculate_days(start, end):
    d1 = datetime.strptime(start, "%d.%m.%Y")
    d2 = datetime.strptime(end, "%d.%m.%Y")
    if d2 < d1:
        raise ValueError("End date cannot be earlier than start date")
    return (d2 - d1).days + 1

def extend_date(date_str, days=3):
    d = datetime.strptime(date_str, "%d.%m.%Y")
    return (d + timedelta(days=days)).strftime("%d.%m.%Y")

# ==========================================
# 🛣️ Road types
# ==========================================
def choose_road_types():
    options = ["Asphalt", "Gravel", "Dirt", "Off-road"]
    print("\nTypes of roads:")
    for i, o in enumerate(options, 1):
        print(f"{i}. {o}")

    choice = input("Choose numbers (comma separated, example: 1,3): ").strip()
    result = []

    for c in choice.split(","):
        if c.strip().isdigit():
            idx = int(c.strip())
            if 1 <= idx <= len(options):
                result.append(options[idx - 1])

    return ", ".join(result) if result else "Asphalt"

# ==========================================
# 🌍 Additional countries
# ==========================================
def choose_additional_countries():
    options = {
        1: "Kyrgyzstan",
        2: "Uzbekistan",
        3: "Tajikistan"
    }

    print("\nAdditional countries:")
    for i, name in options.items():
        print(f"{i}. {name}")

    choice = input("Choose numbers (comma separated) or press Enter to skip: ").strip()
    result = []

    for c in choice.split(","):
        if c.strip().isdigit() and int(c.strip()) in options:
            result.append(options[int(c.strip())])

    return result

# ==========================================
# 🚀 MAIN
# ==========================================
if __name__ == "__main__":

    cars = load_cars()
    selected_cars = choose_cars(cars)

    client_name = input("Client full name: ")
    date_of_birth = input_date("Date of birth")
    address = input("Home address: ")
    phone = input("Phone number: ")
    email = input("Email address: ")

    passport_number = input("Passport / ID number: ")
    passport_issue_date = input_date("Passport issue date")
    passport_issue_by = input("Issued by (authority): ")
    license_num = input("Driver license number: ")

    start_date = input_date("Rental start date")
    end_date = input_date("Rental end date")

    rental_rate = input_float("Daily rental price (USD): ")
    
    try:
        days = calculate_days(start_date, end_date)
    except ValueError as e:
        print(f"❌ {e}")
        exit()

    total = rental_rate * days
    deposit = input_float("Security deposit (USD): ")

    # ===== ADDITIONAL DRIVERS =====
    driver_blocks = {}
    driver_names = {}
    driver_licenses = {}

    for i in range(1, 4):
        driver_blocks[f"{{{{DRIVER{i}_BLOCK}}}}"] = ""
        driver_names[f"{{{{DRIVER{i}_NAME}}}}"] = ""
        driver_licenses[f"{{{{DRIVER{i}_LICENSE}}}}"] = ""

    if input("Add additional drivers? (yes/no): ").lower() == "yes":
        count = int(input("How many drivers (1–3): "))
        for i in range(count):
            name = input("Driver full name: ")
            passport = input("Passport / ID number: ")
            issue_date = input_date("Passport issue date")
            issue_by = input("Issued by: ")
            license_d = input("Driver license number: ")

            driver_names[f"{{{{DRIVER{i+1}_NAME}}}}"] = name
            driver_licenses[f"{{{{DRIVER{i+1}_LICENSE}}}}"] = f"№{license_d}"

            driver_blocks[f"{{{{DRIVER{i+1}_BLOCK}}}}"] = (
                f"{name}, паспорт (удостоверение) №{passport} от {issue_date}, "
                f"выдано {issue_by}; водительские права №{license_d}"
            )

    road_types = choose_road_types()
    extra_countries = choose_additional_countries()

    allowed_countries = "Kazakhstan"
    if extra_countries:
        allowed_countries += ", " + ", ".join(extra_countries)

    outside_kz_block = (
        f"and outside Kazakhstan ({', '.join(extra_countries)})"
        if extra_countries else ""
    )

    contract_number = load_contract_number()
    save_contract_number(contract_number)

    # 🔁 ГЕНЕРАЦИЯ ДЛЯ КАЖДОЙ МАШИНЫ
    for car in selected_cars:

        base_data = {
            "{{CONTRACT_DATE}}": datetime.now().strftime("%d.%m.%Y"),
            "{{CONTRACT_NUMBER}}": contract_number,
            "{{CLIENT_NAME}}": client_name,
            "{{DATE_OF_BIRTH}}": date_of_birth,
            "{{ADDRESS}}": address,
            "{{PHONE}}": phone,
            "{{EMAIL}}": email,
            "{{PASSPORT_NUMBER}}": passport_number,
            "{{PASSPORT_ISSUE_DATE}}": passport_issue_date,
            "{{PASSPORT_ISSUE_BY}}": passport_issue_by,
            "{{DRIVER_LICENSE}}": license_num,

            "{{RENTAL_START}}": start_date,

            "{{RENTAL_RATE}}": f"{rental_rate:.2f}",
            "{{TOTAL_AMOUNT}}": f"{total:.2f}",
            "{{SECURITY_DEPOSIT}}": f"{deposit:.2f}",

            "{{ALLOWED_COUNTRIES}}": allowed_countries,
            "{{TYPES_OF_ROADS}}": road_types,
            "{{OUTSIDE_KZ_BLOCK}}": outside_kz_block,

            "{{CAR_MAKE}}": car["make"],
            "{{CAR_MODEL}}": car["model"],
            "{{CAR_NAME}}": f"{car['make']} {car['model']}",
            "{{CAR_YEAR}}": car["year"],
            "{{CAR_COLOR}}": car["color"],
            "{{CAR_PLATE}}": car["plate"],
            "{{CAR_VIN}}": car["vin"]
        }

        base_data.update(driver_blocks)
        base_data.update(driver_names)
        base_data.update(driver_licenses)

        generate_docs(base_data, f"{client_name.replace(' ', '_')}_{car['plate']}", end_date)

    print("\nAll documents are created successfully.")
    input("Press Enter to exit...")